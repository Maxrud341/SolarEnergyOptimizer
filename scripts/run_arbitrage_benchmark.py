import os
import sys
from pathlib import Path
sys.path.insert(0, str(Path.cwd()))

import pandas as pd
import warnings
warnings.filterwarnings('ignore')

from src.optimizer_v3 import (
    CONFIG_V3,
    run_v2_baseline_v3,
    run_naive_baseline_v3,
    compute_financial_kpis_v3,
)
from src.optimizer_arbitrage import (
    CONFIG_ARBITRAGE,
    run_arbitrage_backtest,
    compute_arbitrage_kpis,
)

def save_styled_excel(res_df: pd.DataFrame, output_file: Path, lang: str = 'ru'):
    """Save formatted Excel workbook with corporate styles in Russian or Czech."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active

    if lang == 'cz':
        ws.title = 'B2G Arbitráž Souhrn'
        title_text = 'Komerční spotová arbitráž FVE (Battery-to-Grid) vs Základní strategie (Naive)'
        rename_dict = {
            'Month': 'Měsíc',
            'Hours': 'Počet hodin',
            'PV_kWh': 'Výroba FVE (kWh)',
            'Load_kWh': 'Spotřeba domu (kWh)',
            'Net_Naive_CZK': 'Fin. výsledek Naive (CZK)',
            'Net_Arbitrage_CZK': 'Fin. výsledek Arbitrage B2G (CZK)',
            'Arbitrage_vs_Naive_CZK': 'Přínos B2G vs Naive (CZK)',
            'Export_Sold_kWh': 'Export do sítě prodáno (kWh)',
            'Curtailment_kWh': 'Omezení přetoků (kWh)',
            'Battery_Cycles': 'Cykly baterie',
        }
        highlight_header = 'Přínos B2G vs Naive (CZK)'
        total_text = 'CELKEM'
    else:
        ws.title = 'B2G Arbitrage Сводка'
        title_text = 'Коммерческий биржевой арбитраж СЭС (Battery-to-Grid) vs Базовая стратегия (Naive)'
        rename_dict = {
            'Month': 'Месяц',
            'Hours': 'Часов в выборке',
            'PV_kWh': 'Генерация СЭС (кВт⋅ч)',
            'Load_kWh': 'Потребление дома (кВт⋅ч)',
            'Net_Naive_CZK': 'Фин. итог Naive (CZK)',
            'Net_Arbitrage_CZK': 'Фин. итог Arbitrage B2G (CZK)',
            'Arbitrage_vs_Naive_CZK': 'Выгода B2G vs Naive (CZK)',
            'Export_Sold_kWh': 'Экспорт в сеть продано (кВт⋅ч)',
            'Curtailment_kWh': 'Ограничение перетоков (кВт⋅ч)',
            'Battery_Cycles': 'Циклов АКБ',
        }
        highlight_header = 'Выгода B2G vs Naive (CZK)'
        total_text = 'ИТОГО'

    # Title
    ws.merge_cells('A1:J1')
    ws['A1'] = title_text
    ws['A1'].font = Font(name='Segoe UI', size=14, bold=True, color='1F497D')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    df_export = res_df.rename(columns=rename_dict)
    headers = list(df_export.columns)

    header_fill = PatternFill(start_color='1F497D', end_color='1F497D', fill_type='solid')
    header_font = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
    ws.row_dimensions[3].height = 30

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    data_font = Font(name='Segoe UI', size=10)
    green_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
    dark_green_font = Font(name='Segoe UI', size=10, bold=True, color='276A3C')

    row_start = 4
    for r_idx, row in df_export.iterrows():
        curr_row = row_start + r_idx
        ws.row_dimensions[curr_row].height = 20
        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(row=curr_row, column=c_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if c_idx in [1, 2]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='right', vertical='center')

            if headers[c_idx-1] == highlight_header:
                cell.fill = green_fill
                cell.font = dark_green_font

    # Total Row
    tot_row = row_start + len(df_export)
    ws.row_dimensions[tot_row].height = 24
    tot_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
    tot_font = Font(name='Segoe UI', size=10, bold=True)
    double_border = Border(
        top=Side(style='double', color='1F497D'),
        bottom=Side(style='double', color='1F497D')
    )
    ws.cell(row=tot_row, column=1, value=total_text).alignment = Alignment(horizontal='center', vertical='center')
    ws.cell(row=tot_row, column=1).font = tot_font
    ws.cell(row=tot_row, column=1).fill = tot_fill

    for c in range(1, len(headers) + 1):
        col_let = get_column_letter(c)
        cell = ws.cell(row=tot_row, column=c)
        cell.font = tot_font
        cell.fill = tot_fill
        cell.border = double_border
        if c in range(2, len(headers) + 1):
            cell.value = f'=SUM({col_let}4:{col_let}{tot_row-1})'
            cell.alignment = Alignment(horizontal='right', vertical='center')

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['G'].width = 24

    wb.save(output_file)
    print(f'Saved Excel ({lang.upper()}) to {output_file}')


def run_all_benchmarks():
    df = pd.read_csv('data/combined_market_data.csv', sep=';')
    arb_log = pd.read_csv('trading_log_arbitrage.csv', sep=';')

    df['_dt'] = pd.to_datetime(df['datetime'])
    arb_log['_dt'] = pd.to_datetime(arb_log['datetime'])
    periods = sorted(df['_dt'].dt.to_period('M').unique())

    print("Running fast baseline (Naive) on combined_market_data.csv...")
    log_naive_full = run_naive_baseline_v3(df, cfg=CONFIG_V3)
    log_naive_full['_dt'] = df['_dt'].values

    results = []
    print(f'Calculating monthly benchmark across {len(periods)} periods...')

    for p in periods:
        p_str = str(p)
        sub_df = df[df['_dt'].dt.to_period('M') == p].drop(columns=['_dt']).reset_index(drop=True)
        if len(sub_df) < 50:
            continue

        sub_arb = arb_log[arb_log['_dt'].dt.to_period('M') == p].drop(columns=['_dt']).reset_index(drop=True)
        sub_naive = log_naive_full[log_naive_full['_dt'].dt.to_period('M') == p].drop(columns=['_dt']).reset_index(drop=True)

        kpi_arb = compute_arbitrage_kpis(sub_arb)
        kpi_naive = compute_financial_kpis_v3(sub_naive)

        pv_tot = sub_arb['pv_kwh'].dropna().sum()
        load_tot = sub_arb['load_kwh'].dropna().sum()
        curt_tot = sub_arb['curtailment_kwh'].dropna().sum() if 'curtailment_kwh' in sub_arb.columns else 0.0

        results.append({
            'Month': p_str,
            'Hours': len(sub_df),
            'PV_kWh': round(pv_tot, 1),
            'Load_kWh': round(load_tot, 1),
            'Net_Naive_CZK': round(kpi_naive['net_profit_czk'], 2),
            'Net_Arbitrage_CZK': round(kpi_arb['net_profit_czk'], 2),
            'Arbitrage_vs_Naive_CZK': round(kpi_arb['net_profit_czk'] - kpi_naive['net_profit_czk'], 2),
            'Export_Sold_kWh': round(kpi_arb['total_sell_kwh'], 1),
            'Curtailment_kWh': round(curt_tot, 1),
            'Battery_Cycles': round(kpi_arb['total_cycles'], 1),
        })

    res_df = pd.DataFrame(results)
    
    # Save CSV
    benchmarks_dir = Path("benchmarks")
    benchmarks_dir.mkdir(parents=True, exist_ok=True)
    csv_file = benchmarks_dir / 'monthly_arbitrage_benchmark.csv'
    res_df.to_csv(csv_file, index=False)
    print(f'\nSaved summary to {csv_file}')

    # Save Excel with formatting (Russian and Czech)
    try:
        excel_file = benchmarks_dir / 'monthly_arbitrage_benchmark.xlsx'
        latest_file = benchmarks_dir / 'monthly_arbitrage_benchmark_latest.xlsx'
        excel_cz_file = benchmarks_dir / 'monthly_arbitrage_benchmark_cz.xlsx'

        save_styled_excel(res_df, excel_file, lang='ru')
        save_styled_excel(res_df, latest_file, lang='ru')
        save_styled_excel(res_df, excel_cz_file, lang='cz')

        # Also save Czech CSV
        csv_cz_file = benchmarks_dir / 'monthly_arbitrage_benchmark_cz.csv'
        rename_cz = {
            'Month': 'Měsíc',
            'Hours': 'Počet hodin',
            'PV_kWh': 'Výroba FVE (kWh)',
            'Load_kWh': 'Spotřeba domu (kWh)',
            'Net_Naive_CZK': 'Fin. výsledek Naive (CZK)',
            'Net_Arbitrage_CZK': 'Fin. výsledek Arbitrage B2G (CZK)',
            'Arbitrage_vs_Naive_CZK': 'Přínos B2G vs Naive (CZK)',
            'Export_Sold_kWh': 'Export do sítě prodáno (kWh)',
            'Curtailment_kWh': 'Omezení přetoků (kWh)',
            'Battery_Cycles': 'Cykly baterie',
        }
        res_df.rename(columns=rename_cz).to_csv(csv_cz_file, index=False)
        print(f'Saved Czech CSV to {csv_cz_file}')
    except Exception as e:
        print(f'Excel export note: {e}')

    return res_df

if __name__ == '__main__':
    run_all_benchmarks()
